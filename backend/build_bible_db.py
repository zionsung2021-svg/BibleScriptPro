import io
import urllib.parse
import sqlite3
import os
from typing import Optional
from fastapi import FastAPI, Query, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

app = FastAPI(title="BibleScript API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DB_PATH = "bible.db"

# 1글자 약어와 표준 전체 명칭 상호 매핑
NAME_VARIANTS = {
    "시편": ["시편", "시"],
    "시": ["시편", "시"],
    "창세기": ["창세기", "창"],
    "출애굽기": ["출애굽기", "출"],
    "느헤미야": ["느헤미야", "느"],
    "에스더": ["에스더", "에"],
    "욥기": ["욥기", "욥"],
    "잠언": ["잠언", "잠"],
    "전도서": ["전도서", "전"],
    "이사야": ["이사야", "사"]
}

def get_unit_name(book: str) -> str:
    return "편" if book.strip() in ["시편", "시"] else "장"

def get_bible_verses_from_db(book: str, chapter: int, start_v: Optional[int] = None, end_v: Optional[int] = None):
    if not os.path.exists(DB_PATH):
        raise HTTPException(
            status_code=500,
            detail="데이터베이스(bible.db)가 존재하지 않습니다."
        )

    clean_book = book.strip()
    unit = get_unit_name(clean_book)
    candidates = NAME_VARIANTS.get(clean_book, [clean_book])
    placeholders = ",".join(["?"] * len(candidates))

    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    # 절 범위가 생략된 경우 해당 장의 1절 ~ 마지막 절 자동 계산
    if start_v is None or start_v <= 0:
        start_v = 1
    if end_v is None or end_v <= 0:
        cur.execute(
            f"SELECT MAX(verse) FROM bible WHERE book IN ({placeholders}) AND chapter = ?",
            candidates + [chapter]
        )
        max_row = cur.fetchone()
        if not max_row or max_row[0] is None:
            conn.close()
            raise HTTPException(
                status_code=404,
                detail=f"{clean_book} {chapter}{unit} 본문이 DB에 없습니다."
            )
        end_v = max_row[0]

    cur.execute(
        f"""
        SELECT verse, content FROM bible 
        WHERE book IN ({placeholders}) AND chapter = ? AND verse BETWEEN ? AND ? 
        ORDER BY verse ASC
        """,
        candidates + [chapter, start_v, end_v]
    )
    rows = cur.fetchall()
    conn.close()

    if not rows:
        raise HTTPException(
            status_code=404,
            detail=f"{clean_book} {chapter}{unit} {start_v}~{end_v}절 본문을 찾을 수 없습니다."
        )

    return [{"verse_num": r[0], "text": r[1]} for r in rows], start_v, end_v

@app.get("/api/search")
def search_bible(
    book: str = Query("시편"),
    chapter: int = Query(23),
    start_verse: Optional[int] = Query(None),
    end_verse: Optional[int] = Query(None),
    version: str = Query("개역개정")
):
    unit = get_unit_name(book)
    verses, start_v, end_v = get_bible_verses_from_db(book, chapter, start_verse, end_verse)
    formatted = "\n\n".join([f"[{book} {chapter}{unit} {v['verse_num']}절]\n{v['text']}" for v in verses])
    return {
        "book": book,
        "chapter": chapter,
        "unit": unit,
        "start_verse": start_v,
        "end_verse": end_v,
        "total_verses": len(verses),
        "reference": f"{book} {chapter}{unit} {start_v}~{end_v}절",
        "verses": verses,
        "text": formatted
    }

@app.get("/api/generate-ppt")
def generate_ppt(
    book: str = Query(...),
    chapter: int = Query(..., ge=1),
    start_verse: Optional[int] = Query(None),
    end_verse: Optional[int] = Query(None),
    version: str = Query("개역개정"),
    font_name: str = Query("Pretendard")
):
    unit = get_unit_name(book)
    verses, start_v, end_v = get_bible_verses_from_db(book, chapter, start_verse, end_verse)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    bg_color = RGBColor(0, 32, 96)       # Navy Blue
    title_color = RGBColor(255, 192, 0)  # Gold Yellow
    body_color = RGBColor(255, 255, 255) # Pure White

    for v in verses:
        slide = prs.slides.add_slide(blank_layout)
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color

        text_length = len(v["text"])

        # 1. 상단 중앙 장/절 제목
        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(1.1))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = f"{book} {chapter}{unit} {v['verse_num']}절"
        p_title.font.name = font_name
        p_title.font.size = Pt(46)
        p_title.font.bold = True
        p_title.font.color.rgb = title_color
        p_title.alignment = PP_ALIGN.CENTER

        # 2. 중앙 본문
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(11.733), Inches(4.8))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True
        p_content = tf_content.paragraphs[0]
        p_content.text = v["text"]
        p_content.font.name = font_name

        if text_length > 100:
            p_content.font.size = Pt(36)
        elif text_length > 55:
            p_content.font.size = Pt(42)
        else:
            p_content.font.size = Pt(48)

        p_content.font.bold = True
        p_content.font.color.rgb = body_color
        p_content.alignment = PP_ALIGN.CENTER

    file_stream = io.BytesIO()
    prs.save(file_stream)
    file_stream.seek(0)

    filename = f"{book}_{chapter}{unit}_{start_v}_{end_v}절.pptx"
    encoded_filename = urllib.parse.quote(filename)

    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
        }
    )